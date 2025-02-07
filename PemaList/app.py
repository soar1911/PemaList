from flask import Flask, request, jsonify, render_template, send_file
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.shared import Length
import os
from datetime import datetime

app = Flask(__name__)

# 確保輸出目錄存在
OUTPUT_DIR = 'output_files'
if not os.path.exists(OUTPUT_DIR):
    os.makedirs(OUTPUT_DIR)

def set_document_orientation_and_font(doc, is_landscape=True):
    section = doc.sections[0]
    if is_landscape:
        # 設定為橫式
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    else:
        # 設定為直式
        section.orientation = WD_ORIENT.PORTRAIT
    
    # 設定邊界為 1.27cm
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.27)

    # 設定預設字型大小
    style = doc.styles['Normal']
    font = style.font
    font.size = Pt(12)

def set_paragraph_format(paragraph):
    """設置段落格式：無間距，固定行高16pt"""
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(0)  # 段落後間距為0
    paragraph_format.space_before = Pt(0)  # 段落前間距為0
    paragraph_format.line_spacing = Pt(16)  # 固定行高16pt
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY  # 設置為固定行高

def add_empty_lines(doc, count):
    for _ in range(count):
        paragraph = doc.add_paragraph('')
        set_paragraph_format(paragraph)

def estimate_line_count(text, max_chars_per_line):
    """估算文字會佔用幾行"""
    if not text:
        return 1
    return -(-len(text) // max_chars_per_line)  # 向上取整除法

def create_word_files(df):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_paths = {}
    LINES_PER_PAGE = 30      # 一頁最多30行
    START_LINE = 20          # 從第20行開始
    MAX_CHARS_PER_LINE = 200  # 根據實際範例調整為200個字元

    # 1. 消災牌位 (直式)
    doc1 = Document()
    set_document_orientation_and_font(doc1, is_landscape=False)
    add_empty_lines(doc1, START_LINE - 1)

    current_line = START_LINE
    for _, row in df.iterrows():
        if pd.notna(row['祈福牌位(隨喜)']):
            content_text = str(row['祈福牌位(隨喜)']).replace('\n', ' ')
            content = f"{row['姓名']}\t{content_text}"
            
            # 更準確的行數估算
            name_length = len(row['姓名'])
            content_length = len(content_text)
            total_length = name_length + content_length + 1  # +1 for tab
            estimated_lines = max(1, -(-total_length // MAX_CHARS_PER_LINE))
            
            if current_line + estimated_lines > LINES_PER_PAGE:
                doc1.add_page_break()
                add_empty_lines(doc1, START_LINE - 1)
                current_line = START_LINE

            paragraph = doc1.add_paragraph(content)
            set_paragraph_format(paragraph)
            current_line += estimated_lines

    file_paths['消災牌位'] = os.path.join(OUTPUT_DIR, f'消災牌位_{timestamp}.docx')
    doc1.save(file_paths['消災牌位'])

    # 2. 超薦牌位 (直式)
    doc2 = Document()
    set_document_orientation_and_font(doc2, is_landscape=False)
    add_empty_lines(doc2, START_LINE - 1)

    current_line = START_LINE
    for _, row in df.iterrows():
        if pd.notna(row['超薦牌位(隨喜)']):
            content_text = str(row['超薦牌位(隨喜)']).replace('\n', ' ')
            content = f"{row['姓名']}\t{content_text}"
            
            # 更準確的行數估算
            name_length = len(row['姓名'])
            content_length = len(content_text)
            total_length = name_length + content_length + 1  # +1 for tab
            estimated_lines = max(1, -(-total_length // MAX_CHARS_PER_LINE))
            
            if current_line + estimated_lines > LINES_PER_PAGE:
                doc2.add_page_break()
                add_empty_lines(doc2, START_LINE - 1)
                current_line = START_LINE

            paragraph = doc2.add_paragraph(content)
            set_paragraph_format(paragraph)
            current_line += estimated_lines

    file_paths['超薦牌位'] = os.path.join(OUTPUT_DIR, f'超薦牌位_{timestamp}.docx')
    doc2.save(file_paths['超薦牌位'])


    # 3. 功德主 (保持橫式)
    doc3 = Document()
    set_document_orientation_and_font(doc3, is_landscape=True)
    doc3.add_heading('功德主', 0)

    table = doc3.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    header_cells = table.rows[0].cells
    headers = ['姓名', 'Email', '行動電話', '參贊功德主']

    for i, header in enumerate(headers):
        header_cells[i].text = header

    for _, row in df.iterrows():
        if pd.notna(row['參贊功德主']):
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['姓名'])
            row_cells[1].text = str(row['Email'])
            row_cells[2].text = str(row['行動電話'])
            row_cells[3].text = str(row['參贊功德主'])

    file_paths['功德主'] = os.path.join(OUTPUT_DIR, f'功德主_{timestamp}.docx')
    doc3.save(file_paths['功德主'])

    return file_paths

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/process_excel', methods=['POST'])
def process_excel():
    try:
        # 確保輸出目錄存在
        if not os.path.exists(OUTPUT_DIR):
            os.makedirs(OUTPUT_DIR)

        if 'file' not in request.files:
            return jsonify({'error': '沒有檔案'}), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': '沒有選擇檔案'}), 400

        # 讀取Excel檔案
        df = pd.read_excel(file)

        # 檢查必要欄位
        required_columns = ['姓名', 'Email', '行動電話', '祈福牌位(隨喜)', '超薦牌位(隨喜)', '參贊功德主']
        missing_columns = [col for col in required_columns if col not in df.columns]

        if missing_columns:
            return jsonify({'error': f'缺少欄位：{", ".join(missing_columns)}'}), 400

        # 處理並創建Word檔案
        file_paths = create_word_files(df)

        # 返回檔案名稱（不是完整路徑）
        return jsonify({
            'message': '處理完成',
            'files': {
                'xiazai': os.path.basename(file_paths['消災牌位']),
                'chaojian': os.path.basename(file_paths['超薦牌位']),
                'gongde': os.path.basename(file_paths['功德主'])
            }
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/download/<filename>')
def download_file(filename):
    try:
        if not os.path.exists(OUTPUT_DIR):
            return jsonify({'error': '輸出目錄不存在'}), 400

        file_path = os.path.join(OUTPUT_DIR, filename)
        if not os.path.exists(file_path):
            return jsonify({'error': '檔案不存在'}), 404

        return send_file(
            file_path,
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return str(e), 400

if __name__ == '__main__':
    app.run(debug=True)
